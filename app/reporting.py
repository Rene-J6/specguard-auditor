import io
import re
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt


def _report_filename(document_name):
    stem = Path(document_name or "document").stem
    safe_stem = re.sub(r"[^A-Za-z0-9_-]+", "_", stem).strip("_") or "document"
    return f"{safe_stem}_audit_report.docx"


def _add_labeled_text(document, label, text):
    paragraph = document.add_paragraph()
    paragraph.add_run(f"{label}: ").bold = True
    paragraph.add_run(text or "—")


def build_docx_report(report):
    document = Document()
    normal_style = document.styles["Normal"]
    normal_style.font.name = "Calibri"
    normal_style.font.size = Pt(10)

    title = document.add_heading("SpecGuard SRS Audit Report", level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _add_labeled_text(document, "Document", report["document_name"])
    _add_labeled_text(document, "Generated", report["generated_at"])

    document.add_heading("Audit Summary", level=1)
    document.add_paragraph(
        f"Total requirements: {report['total_requirements']}", style="List Bullet"
    )
    document.add_paragraph(
        f"Distinct standard sources: {len(report['standard_counts'])}",
        style="List Bullet",
    )
    document.add_paragraph(
        f"Expert-adjusted requirements: {report['expert_modified_count']}",
        style="List Bullet",
    )
    if report["standard_counts"]:
        document.add_paragraph("Matched standard/source distribution:")
        for source, count in report["standard_counts"].items():
            document.add_paragraph(f"{source}: {count}", style="List Bullet 2")

    document.add_heading("Detailed Findings", level=1)
    for index, item in enumerate(report["items"], start=1):
        document.add_heading(f"Requirement {index}", level=2)
        _add_labeled_text(document, "Original requirement", item["requirement"])
        _add_labeled_text(document, "Matched standard/clause", item["matched_rule"])
        _add_labeled_text(document, "AI analysis", item["analysis"])
        _add_labeled_text(document, "Suggested rewrite", item["suggested_rewrite"])
        if item["expert_evaluation"] is not None:
            _add_labeled_text(
                document, "Expert evaluation", item["expert_evaluation"]
            )
        if item["expert_rewrite"] is not None:
            _add_labeled_text(document, "Expert rewrite", item["expert_rewrite"])

    output = io.BytesIO()
    document.save(output)
    return output.getvalue(), _report_filename(report["document_name"])
