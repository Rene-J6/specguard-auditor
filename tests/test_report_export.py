import io
import tempfile
import unittest
from pathlib import Path

from docx import Document

from app import database


class ReportExportTests(unittest.TestCase):
    def setUp(self):
        self.temp_dir = tempfile.TemporaryDirectory()
        root = Path(self.temp_dir.name)
        self.original_pool_path = database.SQLITE_DB_PATH
        self.original_dataset_path = database.EXPORT_JSONL_PATH
        database.SQLITE_DB_PATH = str(root / "training_pool.db")
        database.EXPORT_JSONL_PATH = str(root / "dataset.jsonl")
        Path(database.EXPORT_JSONL_PATH).write_text("", encoding="utf-8")
        database.init_sqlite_pool()

    def tearDown(self):
        database.SQLITE_DB_PATH = self.original_pool_path
        database.EXPORT_JSONL_PATH = self.original_dataset_path
        self.temp_dir.cleanup()

    def add_approved_document(self):
        database.log_audit_to_db(
            "owner",
            "The system shall process reports quickly.",
            "ISO29148.pdf | Requirements shall be measurable.",
            "### Analysis\nThe timing is ambiguous.\n\n"
            "### Suggested Rewrite\nThe system shall process a report within 2 seconds.",
            "doc_report",
            "Customer SRS?.pdf",
        )
        database.lock_requirement_set("doc_report", "expert")
        requirement = database.fetch_document_requirements(
            "doc_report", "expert"
        )["requirements"][0]
        database.submit_expert_review(
            "doc_report",
            "expert",
            [{
                "id": requirement["id"],
                "expert_evaluation": "Define a measurable response time.",
                "expert_rewrite": "The system shall process a report within 1 second.",
            }],
            True,
        )

    def document_text(self, content):
        document = Document(io.BytesIO(content))
        blocks = [paragraph.text for paragraph in document.paragraphs]
        for table in document.tables:
            for row in table.rows:
                blocks.extend(cell.text for cell in row.cells)
        return "\n".join(blocks)

    def test_report_data_is_owner_scoped_and_hides_internal_metadata(self):
        self.add_approved_document()

        report = database.fetch_user_document_report("doc_report", "owner")
        forbidden = database.fetch_user_document_report("doc_report", "other")

        self.assertIsNone(forbidden)
        self.assertEqual(report["document_name"], "Customer SRS?.pdf")
        self.assertEqual(report["total_requirements"], 1)
        self.assertEqual(report["standard_counts"], {"ISO29148.pdf": 1})
        self.assertEqual(report["expert_modified_count"], 1)
        self.assertNotIn("lifecycle_status", report)
        self.assertNotIn("assigned_expert", report)
        self.assertNotIn("export_status", report)

    def test_docx_contains_audit_fields_without_internal_status(self):
        from app.reporting import build_docx_report

        self.add_approved_document()
        report = database.fetch_user_document_report("doc_report", "owner")

        content, filename = build_docx_report(report)
        text = self.document_text(content)

        self.assertEqual(filename, "Customer_SRS_audit_report.docx")
        self.assertIn("SpecGuard SRS Audit Report", text)
        self.assertIn("The timing is ambiguous.", text)
        self.assertIn("The system shall process a report within 2 seconds.", text)
        self.assertIn("Define a measurable response time.", text)
        self.assertIn("The system shall process a report within 1 second.", text)
        self.assertNotIn("PENDING", text)
        self.assertNotIn("APPROVED", text)
        self.assertNotIn("assigned_expert", text)
        self.assertNotIn("export_status", text)

    def test_rejected_expert_content_is_not_exposed(self):
        database.log_audit_to_db(
            "owner", "The system shall work.", "Rule", "### Analysis\nAI issue",
            "doc_rejected", "rejected.pdf",
        )
        database.lock_requirement_set("doc_rejected", "expert")
        requirement = database.fetch_document_requirements(
            "doc_rejected", "expert"
        )["requirements"][0]
        database.submit_expert_review(
            "doc_rejected",
            "expert",
            [{
                "id": requirement["id"],
                "expert_evaluation": "Private rejected evaluation",
                "expert_rewrite": "Private rejected rewrite",
            }],
            False,
        )

        report = database.fetch_user_document_report("doc_rejected", "owner")

        self.assertIsNone(report["items"][0]["expert_evaluation"])
        self.assertIsNone(report["items"][0]["expert_rewrite"])


if __name__ == "__main__":
    unittest.main()
